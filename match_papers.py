import docx, re, difflib, json

sched_doc = docx.Document('../ISPSW_NC_2026_Programme_and_Paper_Presentation_Schedule.docx')
sched_papers = []
for t_idx in [7, 9, 11, 13, 15, 17, 19, 21, 23]:
    t = sched_doc.tables[t_idx]
    for r in t.rows[1:]:
        cells = [c.text.strip().replace('\n', ' ') for c in r.cells]
        if len(cells) >= 3 and cells[0]:
            sched_papers.append({'code': cells[0].upper(), 'title': cells[1], 'authors': cells[2], 't_idx': t_idx})

abs_doc = docx.Document(r'C:\Users\UE\.gemini\antigravity\brain\537e12af-88c3-46c0-86a7-2aaa7fa4218f\.user_uploaded\media_1791267294499.docx')
abs_blocks = {}
current_code = None
current_paras = []
for p in docx.Document(r'C:\Users\UE\.gemini\antigravity\brain\537e12af-88c3-46c0-86a7-2aaa7fa4218f\.user_uploaded\media_1791267294499.docx').paragraphs:
    t = p.text.strip()
    if not t: continue
    m = re.search(r'ABSTRACT CODE\s*:\s*([A-Za-z0-9\-]+)', t)
    if m:
        if current_code: abs_blocks[current_code] = current_paras
        current_code = m.group(1).upper()
        current_paras = [t]
    else:
        if current_code: current_paras.append(t)
if current_code: abs_blocks[current_code] = current_paras

print(f"Sched papers: {len(sched_papers)}, Abs blocks: {len(abs_blocks)}")

matched_by_code = 0
unmatched_sched = []
for sp in sched_papers:
    if sp['code'] in abs_blocks:
        matched_by_code += 1
    else:
        unmatched_sched.append(sp)

print(f"Matched directly by code: {matched_by_code}, Unmatched: {len(unmatched_sched)}")
for sp in unmatched_sched:
    best_ratio = 0
    best_code = None
    best_title = ''
    for acode, aparas in abs_blocks.items():
        atitle = aparas[1] if len(aparas) > 1 else ''
        ratio = difflib.SequenceMatcher(None, sp['title'].lower(), atitle.lower()).ratio()
        if ratio > best_ratio:
            best_ratio = ratio
            best_code = acode
            best_title = atitle
    print(f"Sched {sp['code']}: '{sp['title'][:40]}' -> Best Abs {best_code} ({best_ratio:.2f}): '{best_title[:40]}'")
