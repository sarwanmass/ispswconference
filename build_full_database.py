import docx
import json
import re
import base64
import os

print("Assembling comprehensive conference database with all scheduled & accepted papers...")

# 1. Base64 logos
ispsw_b64 = ""
davangere_b64 = ""
if os.path.exists("assets_schedule/img_0.png"):
    with open("assets_schedule/img_0.png", "rb") as f:
        ispsw_b64 = "data:image/png;base64," + base64.b64encode(f.read()).decode("utf-8")

if os.path.exists("assets_schedule/img_2.png"):
    with open("assets_schedule/img_2.png", "rb") as f:
        davangere_b64 = "data:image/png;base64," + base64.b64encode(f.read()).decode("utf-8")

# 2. Parse abstracts docx
abs_path = r"C:\Users\UE\.gemini\antigravity\brain\537e12af-88c3-46c0-86a7-2aaa7fa4218f\.user_uploaded\media_1791267294499.docx"
abs_doc = docx.Document(abs_path)

abs_blocks = {}
curr_code = None
curr_lines = []

for p in abs_doc.paragraphs:
    t = p.text.strip()
    if not t:
        continue
    m = re.search(r'ABSTRACT CODE\s*:\s*([A-Za-z0-9\-–_—]+)', t)
    if m:
        if curr_code:
            abs_blocks[curr_code] = curr_lines
        curr_code = m.group(1).strip().upper()
        curr_lines = [t]
    else:
        if curr_code:
            curr_lines.append(t)

if curr_code:
    abs_blocks[curr_code] = curr_lines

print(f"Total abstract blocks parsed: {len(abs_blocks)}")

def process_abstract_block(code, lines):
    title = ""
    authors_affil = []
    abstract_paras = []
    keywords = []
    
    state = "title"
    for line in lines[1:]:
        l_upper = line.strip().upper()
        if l_upper == "ABSTRACT":
            state = "abstract"
            continue
        if line.strip().lower().startswith("keywords:"):
            kw_match = re.sub(r'(?i)^keywords:\s*', '', line)
            keywords = [k.strip(' .,;') for k in re.split(r'[,;]', kw_match) if k.strip()]
            state = "keywords"
            continue
            
        if state == "title":
            if not title:
                title = line
                state = "authors"
            else:
                title += " " + line
        elif state == "authors":
            authors_affil.append(line)
        elif state == "abstract":
            abstract_paras.append(line)
        elif state == "keywords":
            if "@" in line or "contact" in line.lower() or "author" in line.lower() or "dr." in line.lower() or "ph:" in line.lower():
                authors_affil.append(line)
            else:
                extra_kws = [k.strip(' .,;') for k in re.split(r'[,;]', line) if k.strip()]
                keywords.extend(extra_kws)
                
    return {
        "code": code,
        "title": title.strip(),
        "authors_affil": "\n".join(authors_affil).strip(),
        "abstract": "\n\n".join(abstract_paras).strip(),
        "keywords": keywords
    }

parsed_abstracts = {code: process_abstract_block(code, lines) for code, lines in abs_blocks.items()}

# 3. Parse Schedule docx
sched_path = "../ISPSW_NC_2026_Programme_and_Paper_Presentation_Schedule.docx"
sched_doc = docx.Document(sched_path)

session_configs = [
    {
        "t_meta": 6, "t_papers": 7,
        "session_id": "TS-1-T1",
        "session_num": "I",
        "track_num": "1",
        "session_title": "Technical Session I · Track 1",
        "day": "Day 01",
        "date": "07.10.2026",
        "day_full": "Day 01 · 07.10.2026, Wednesday",
        "time": "4:45 – 5:30 pm",
        "venue": "MBA Auditorium, Davangere University",
        "themes": "Theme 1: Innovative Technology for Learning and Development in Social Work"
    },
    {
        "t_meta": 8, "t_papers": 9,
        "session_id": "TS-1-T2",
        "session_num": "I",
        "track_num": "2",
        "session_title": "Technical Session I · Track 2",
        "day": "Day 01",
        "date": "07.10.2026",
        "day_full": "Day 01 · 07.10.2026, Wednesday",
        "time": "4:45 – 5:30 pm",
        "venue": "MBA Lecture Hall-01",
        "themes": "Theme 1 & Theme 2: AI & Technology for Intervention"
    },
    {
        "t_meta": 10, "t_papers": 11,
        "session_id": "TS-1-T3",
        "session_num": "I",
        "track_num": "3",
        "session_title": "Technical Session I · Track 3",
        "day": "Day 01",
        "date": "07.10.2026",
        "day_full": "Day 01 · 07.10.2026, Wednesday",
        "time": "4:45 – 5:30 pm",
        "venue": "MBA Lecture Hall-02",
        "themes": "Theme 2: Innovation in Technology and AI for Intervention in Social Work"
    },
    {
        "t_meta": 12, "t_papers": 13,
        "session_id": "TS-2-T1",
        "session_num": "II",
        "track_num": "1",
        "session_title": "Technical Session II · Track 1",
        "day": "Day 02",
        "date": "08.10.2026",
        "day_full": "Day 02 · 08.10.2026, Thursday",
        "time": "4:15 – 5:30 pm",
        "venue": "MBA Auditorium, Davangere University",
        "themes": "Theme 2 & Theme 3: Gender Equity and Social Inclusion"
    },
    {
        "t_meta": 14, "t_papers": 15,
        "session_id": "TS-2-T2",
        "session_num": "II",
        "track_num": "2",
        "session_title": "Technical Session II · Track 2",
        "day": "Day 02",
        "date": "08.10.2026",
        "day_full": "Day 02 · 08.10.2026, Thursday",
        "time": "4:15 – 5:30 pm",
        "venue": "MBA Lecture Hall-01",
        "themes": "Theme 3: Innovation in Tech: Gender Equity and Social Inclusion"
    },
    {
        "t_meta": 16, "t_papers": 17,
        "session_id": "TS-2-T3",
        "session_num": "II",
        "track_num": "3",
        "session_title": "Technical Session II · Track 3",
        "day": "Day 02",
        "date": "08.10.2026",
        "day_full": "Day 02 · 08.10.2026, Thursday",
        "time": "4:15 – 5:30 pm",
        "venue": "MBA Lecture Hall-02",
        "themes": "Theme 3: Innovation in Tech: Gender Equity and Social Inclusion"
    },
    {
        "t_meta": 18, "t_papers": 19,
        "session_id": "TS-3-T1",
        "session_num": "III",
        "track_num": "1",
        "session_title": "Technical Session III · Track 1",
        "day": "Day 03",
        "date": "09.10.2026",
        "day_full": "Day 03 · 09.10.2026, Friday",
        "time": "12:30 – 1:15 pm",
        "venue": "MBA Auditorium, Davangere University",
        "themes": "Theme 3: Innovation in Tech: Gender Equity and Social Inclusion"
    },
    {
        "t_meta": 20, "t_papers": 21,
        "session_id": "TS-3-T2",
        "session_num": "III",
        "track_num": "2",
        "session_title": "Technical Session III · Track 2",
        "day": "Day 03",
        "date": "09.10.2026",
        "day_full": "Day 03 · 09.10.2026, Friday",
        "time": "12:30 – 1:15 pm",
        "venue": "MBA Lecture Hall-01",
        "themes": "Theme 3 & Theme 4: Technology Adoption – Policy, Ethics, and Regulation"
    },
    {
        "t_meta": 22, "t_papers": 23,
        "session_id": "TS-3-T3",
        "session_num": "III",
        "track_num": "3",
        "session_title": "Technical Session III · Track 3",
        "day": "Day 03",
        "date": "09.10.2026",
        "day_full": "Day 03 · 09.10.2026, Friday",
        "time": "12:30 – 1:15 pm",
        "venue": "MBA Lecture Hall-02",
        "themes": "Theme 4 & Theme 5: Technologies for Viksit Bharat 2047"
    }
]

alias_map = {
    "F2": "G9",
    "G7": "F9",
    "F10": "J20",
    "H1": "A5",
    "D11": "B6",
    "D9": "C8",
    "D8": "A2",
    "G10": "E6",
    "J8": "F11"
}

sessions_data = []
all_papers = []
matched_abstract_codes = set()

for cfg in session_configs:
    meta_table = sched_doc.tables[cfg["t_meta"]]
    meta_raw = " ".join([c.text.strip().replace("\n", " ") for r in meta_table.rows for c in r.cells])
    
    chairperson = ""
    discussant = ""
    rapporteurs = ""
    
    if "Chairperson" in meta_raw:
        p1 = meta_raw.split("Chairperson")[1]
        if "Discussant" in p1:
            parts = p1.split("Discussant")
            chairperson = parts[0].strip(" :|")
            p2 = parts[1]
            if "Rapporteurs" in p2:
                parts2 = p2.split("Rapporteurs")
                discussant = parts2[0].strip(" :|")
                rapporteurs = parts2[1].replace("& Session In charge", "").strip(" :|")
            else:
                discussant = p2.strip(" :|")
        else:
            chairperson = p1.strip(" :|")
            
    discussant = re.sub(r'\|\s*Discussant.*$', '', discussant).strip(" :|")
    
    sess_obj = {
        "session_id": cfg["session_id"],
        "session_num": cfg["session_num"],
        "track_num": cfg["track_num"],
        "session_title": cfg["session_title"],
        "day": cfg["day"],
        "date": cfg["date"],
        "day_full": cfg["day_full"],
        "time": cfg["time"],
        "venue": cfg["venue"],
        "themes": cfg["themes"],
        "chairperson": chairperson,
        "discussant": discussant,
        "rapporteurs": rapporteurs,
        "papers": []
    }
    
    p_table = sched_doc.tables[cfg["t_papers"]]
    for r in p_table.rows[1:]:
        cells = [c.text.strip().replace("\n", " ") for c in r.cells]
        if len(cells) >= 3 and cells[0]:
            code = cells[0].strip().upper()
            title = cells[1].strip()
            authors = cells[2].strip()
            remarks = cells[3].strip() if len(cells) > 3 else ""
            
            abs_info = None
            if code in parsed_abstracts:
                abs_info = parsed_abstracts[code]
                matched_abstract_codes.add(code)
            elif code in alias_map and alias_map[code] in parsed_abstracts:
                abs_info = parsed_abstracts[alias_map[code]]
                matched_abstract_codes.add(alias_map[code])
            else:
                for acode, aobj in parsed_abstracts.items():
                    if aobj["title"] and (aobj["title"][:22].lower() in title.lower() or title[:22].lower() in aobj["title"].lower()):
                        abs_info = aobj
                        matched_abstract_codes.add(acode)
                        break
                        
            abstract_text = abs_info["abstract"] if abs_info else ""
            keywords_list = abs_info["keywords"] if abs_info else []
            affil_info = abs_info["authors_affil"] if abs_info else ""
            full_title = abs_info["title"] if abs_info and len(abs_info["title"]) > len(title) else title
            
            # Map theme cleanly
            theme_name = remarks if (remarks and remarks != "—") else cfg["themes"]
            if not theme_name or "theme" not in theme_name.lower():
                if code.startswith("A") or code == "B1":
                    theme_name = "Theme 1: Innovative Technology for Learning and Development in Social Work"
                elif code.startswith("B") or code.startswith("C") or (code.startswith("D") and int(code[1:]) <= 12):
                    theme_name = "Theme 2: Innovation in Technology and AI for Intervention in Social Work"
                elif code.startswith("E") or code.startswith("F") or code.startswith("G") or code in ["D13", "D14", "H2", "H3"]:
                    theme_name = "Theme 3: Innovation in Tech: Gender Equity and Social Inclusion"
                elif code.startswith("H") or code in ["J1", "J2", "J3"]:
                    theme_name = "Theme 4: Technology Adoption – Policy, Ethics, and Regulation"
                elif code.startswith("J"):
                    theme_name = "Theme 5: Technologies for Viksit Bharat 2047"
                else:
                    theme_name = cfg["themes"]

            paper_item = {
                "code": code,
                "title": full_title,
                "authors": authors,
                "affiliation": affil_info,
                "abstract": abstract_text,
                "keywords": keywords_list,
                "theme": theme_name,
                "session_id": cfg["session_id"],
                "session_title": cfg["session_title"],
                "day": cfg["day"],
                "date": cfg["date"],
                "day_full": cfg["day_full"],
                "time": cfg["time"],
                "venue": cfg["venue"],
                "chairperson": chairperson,
                "discussant": discussant,
                "rapporteurs": rapporteurs,
                "is_scheduled": True
            }
            sess_obj["papers"].append(paper_item)
            all_papers.append(paper_item)
            
    sessions_data.append(sess_obj)

# Now also add supplementary / accepted research papers from abstract book not matched directly to a schedule slot
supplementary_papers = []
theme_mapping_for_supp = {
    "J19": "Theme 1: Innovative Technology for Learning and Development in Social Work",
    "J15": "Theme 2: Innovation in Technology and AI for Intervention in Social Work",
    "J16": "Theme 2: Innovation in Technology and AI for Intervention in Social Work",
    "J17": "Theme 2: Innovation in Technology and AI for Intervention in Social Work",
    "J20": "Theme 2: Innovation in Technology and AI for Intervention in Social Work",
    "J24": "Theme 2: Innovation in Technology and AI for Intervention in Social Work",
    "—": "Theme 2: Innovation in Technology and AI for Intervention in Social Work",
    "J10": "Theme 3: Innovation in Tech: Gender Equity and Social Inclusion",
    "J11": "Theme 3: Innovation in Tech: Gender Equity and Social Inclusion",
    "J12": "Theme 3: Innovation in Tech: Gender Equity and Social Inclusion",
    "J13": "Theme 3: Innovation in Tech: Gender Equity and Social Inclusion",
    "J14": "Theme 3: Innovation in Tech: Gender Equity and Social Inclusion",
    "J18": "Theme 5: Technologies for Viksit Bharat 2047",
    "J22": "Theme 5: Technologies for Viksit Bharat 2047",
    "J23": "Theme 5: Technologies for Viksit Bharat 2047"
}

for acode, aobj in parsed_abstracts.items():
    if acode not in matched_abstract_codes and acode not in ["A1", "A2", "A3", "A4", "A5", "A6", "A7", "A8", "A9", "A10", "B1"]:
        # Find theme
        theme_str = theme_mapping_for_supp.get(acode, "Theme 2: Innovation in Technology and AI for Intervention in Social Work")
        clean_authors = aobj["authors_affil"].split("\n")[0] if aobj["authors_affil"] else "Faculty / Scholar"
        supp_item = {
            "code": acode if acode != "—" else "ABS-DMF",
            "title": aobj["title"],
            "authors": clean_authors,
            "affiliation": aobj["authors_affil"],
            "abstract": aobj["abstract"],
            "keywords": aobj["keywords"],
            "theme": theme_str,
            "session_id": "TS-SUPP",
            "session_title": "Accepted Research Paper / Abstract Volume",
            "day": "Day 02 / Day 03",
            "date": "08-09.10.2026",
            "day_full": "Accepted Academic Abstract · ISPSW 2026",
            "time": "Research Presentations Track",
            "venue": "MBA Academic Complex, Davangere University",
            "chairperson": "Conference Scientific Committee",
            "discussant": "Editorial & Advisory Board",
            "rapporteurs": "DoS in Social Work, Davangere University",
            "is_scheduled": False
        }
        all_papers.append(supp_item)
        supplementary_papers.append(supp_item)

print(f"Total scheduled papers: {len(all_papers) - len(supplementary_papers)}")
print(f"Total supplementary accepted abstracts: {len(supplementary_papers)}")
print(f"Total papers & abstracts in master database: {len(all_papers)}")

# Program Days Events
def get_clean_table_events(table):
    events = []
    for r in table.rows[1:]:
        cells = [c.text.strip().replace("\n", " ") for c in r.cells]
        if len(cells) >= 4 and cells[1]:
            sl = cells[0].strip()
            time = cells[1].strip()
            event = cells[2].strip()
            details = cells[3].strip() if len(cells) > 3 else ""
            venue = cells[4].strip() if len(cells) > 4 else ""
            
            cat = "General"
            ev_lower = (event + " " + details).lower()
            if "inauguration" in ev_lower or "welcome" in ev_lower or "valedictory" in ev_lower or "presidential" in ev_lower or "invocation" in ev_lower:
                cat = "Ceremonial"
            elif "keynote" in ev_lower or "plenary" in ev_lower or "special address" in ev_lower:
                cat = "Plenary"
            elif "panel" in ev_lower:
                cat = "Panel"
            elif "technical session" in ev_lower or "paper presentation" in ev_lower:
                cat = "Technical"
            elif "award" in ev_lower or "poster" in ev_lower or "workshop" in ev_lower or "agm" in ev_lower or "general body" in ev_lower:
                cat = "Special"
            elif "cultural" in ev_lower:
                cat = "Cultural"
            elif "lunch" in ev_lower or "dinner" in ev_lower or "tea" in ev_lower:
                cat = "Networking / Meals"
                
            events.append({
                "sl": sl,
                "time": time,
                "event": event,
                "details": details,
                "venue": venue,
                "category": cat
            })
    return events

day1_events = get_clean_table_events(sched_doc.tables[2])
day2_events = get_clean_table_events(sched_doc.tables[3])
day3_events = get_clean_table_events(sched_doc.tables[4])

committee_contacts = [
    {"role": "Registration & Helpdesk", "name": "Dr. Thippesh K.", "phone": "+91 98804 52671", "email": "dumswispswnc2026@gmail.com", "desig": "Assistant Professor & Organizing Secretary"},
    {"role": "Accommodation & Transportation", "name": "Dr. Shivalingappa B. P.", "phone": "+91 98861 41887", "email": "bpshivumsw@gmail.com", "desig": "Professor & Chairman, DoS in Social Work"},
    {"role": "Programmes & Sessions", "name": "Dr. Lokesh M. U.", "phone": "+91 99455 02607", "email": "lokeshmutut@gmail.com", "desig": "Professor, DoS in Social Work"},
    {"role": "Publication & Technical", "name": "Dr. Pradeep B. S.", "phone": "+91 73385 84639", "email": "pradeepwagonr@gmail.com", "desig": "Professor, DoS in Social Work"},
    {"role": "Food & Cultural Committee", "name": "Dr. Patwardhan Rathod", "phone": "+91 80737 67894", "email": "prathod1970@gmail.com", "desig": "Assistant Professor, DoS in Social Work"}
]

dignitaries = [
    {
        "name": "Prof. B. E. Rangaswamy",
        "title": "Hon’ble Vice-Chancellor",
        "org": "Davangere University, Shivagangotri Campus",
        "role": "Chief Patron & Presidential Address",
        "tag": "Chief Patron"
    },
    {
        "name": "Prof. Ramesh B.",
        "title": "Hon’ble Vice-Chancellor & President, ISPSW",
        "org": "Dr. Manmohan Singh Bengaluru City University & ISPSW",
        "role": "Guest of Honour & Conference President",
        "tag": "Conference President"
    },
    {
        "name": "Dr. Harish Hande",
        "title": "Renowned Social Entrepreneur & Ramon Magsaysay Awardee",
        "org": "CEO & Founder, SELCO India Pvt. Ltd., Bengaluru",
        "role": "Keynote Speaker (Day 1: 11:40 am – 12:30 pm)",
        "tag": "Keynote Speaker"
    },
    {
        "name": "Dr. R. Dhanasekara Pandian",
        "title": "Professor, Dept. of Psychiatric Social Work",
        "org": "NIMHANS, Bengaluru",
        "role": "General Secretary, ISPSW (Vote of Thanks)",
        "tag": "General Secretary ISPSW"
    },
    {
        "name": "Dr. I. A. Sharif",
        "title": "Former Professor & Head, Dept. of Psychiatric Social Work",
        "org": "NIMHANS, Bengaluru",
        "role": "Special Address: New Avenues for Social Work Research (Day 2)",
        "tag": "Eminent Speaker"
    },
    {
        "name": "Dr. Anish V. Cherian",
        "title": "Additional Professor, Dept. of Psychiatric Social Work",
        "org": "NIMHANS, Bengaluru",
        "role": "Workshop Lead: Suicide Prevention (Day 2: 3:30 – 4:15 pm)",
        "tag": "Workshop Lead"
    }
]

themes_list = [
    {"id": "Theme 1", "key": "Theme 1", "title": "Innovative Technology for Learning and Development in Social Work", "icon": "fa-graduation-cap", "color": "#1e40af", "badge": "Learning & Tech", "desc": "E-learning, digital tools, educational technology, curriculum integration, vernacular AI and capacity building for social work professionals."},
    {"id": "Theme 2", "key": "Theme 2", "title": "Innovation in Technology and AI for Intervention in Social Work", "icon": "fa-brain", "color": "#6d28d9", "badge": "AI & Interventions", "desc": "Artificial intelligence, machine learning, predictive analytics, wearable health monitors, digital CSR, and psychosocial interventions."},
    {"id": "Theme 3", "key": "Theme 3", "title": "Innovation in Tech: Gender Equity and Social Inclusion", "icon": "fa-people-roof", "color": "#047857", "badge": "Gender & Inclusion", "desc": "Assistive technologies, tribal empowerment, women and child welfare, disability inclusion, and bridging regional digital divides."},
    {"id": "Theme 4", "key": "Theme 4", "title": "Technology Adoption – Policy, Ethics, and Regulation", "icon": "fa-scale-balanced", "color": "#b45309", "badge": "Policy & Ethics", "desc": "Data protection, algorithmic accountability, NCAHP Act, labour compliance, digital surveillance, and ethical governance in human services."},
    {"id": "Theme 5", "key": "Theme 5", "title": "Technologies for Viksit Bharat 2047", "icon": "fa-landmark", "color": "#b91c1c", "badge": "Viksit Bharat 2047", "desc": "National digital missions, rural transformation, farmer empowerment, health informatics, and societal progress toward Viksit Bharat 2047."}
]

final_database = {
    "conference_meta": {
        "title": "ANNUAL NATIONAL CONFERENCE OF ISPSW – 2026",
        "sub_title": "Innovative Technologies for Social Work Practice, Research and Development",
        "organizers": "Department of Studies in Social Work, Davangere University & Indian Society of Professional Social Work (ISPSW)",
        "dates": "7th – 9th OCTOBER, 2026",
        "dates_short": "Oct 7–9, 2026",
        "venue": "MBA Auditorium & Academic Complex, Shivagangotri Campus, Davangere University, Tholahunase, Davanagere – 577007, Karnataka, India",
        "email": "dumswispswnc2026@gmail.com",
        "ispsw_logo": ispsw_b64,
        "davangere_logo": davangere_b64,
        "total_papers": len(all_papers),
        "total_sessions": len(sessions_data)
    },
    "themes": themes_list,
    "dignitaries": dignitaries,
    "committee_contacts": committee_contacts,
    "programme_schedule": {
        "day1": {"date": "07.10.2026", "day": "Wednesday", "title": "Day 01 · Inauguration, Plenary I, Panel I & Technical Session I", "events": day1_events},
        "day2": {"date": "08.10.2026", "day": "Thursday", "title": "Day 02 · Plenaries II & III, Panel II, Special Address, Workshop & Technical Session II", "events": day2_events},
        "day3": {"date": "09.10.2026", "day": "Friday", "title": "Day 03 · Plenary IV, Panel III, Technical Session III & Valedictory Ceremony", "events": day3_events}
    },
    "technical_sessions": sessions_data,
    "papers": all_papers
}

with open("conference_database.json", "w", encoding="utf-8") as f:
    json.dump(final_database, f, ensure_ascii=False, indent=2)

print("Database written to conference_database.json successfully!")
